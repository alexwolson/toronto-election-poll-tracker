"""Build Chow's 2023 ward benchmark for the results Map (live-projection #108).

Input, beside this script:
  2023_office_of_the_mayor.xlsx  the City's certified poll-by-poll results of the 2023 by-election
                                 for mayor, unmodified from the open data portal (sha256 below):
                                 one sheet per City ward, each candidate's votes by subdivision
                                 (election day, advance, mail, long-term care) and a ward total

Each ward's share is Chow's ward total over every candidate's, all voting channels. The 2023
by-election used the same 25-ward map as 2026. Writes src/data/chow-2023.json, shares in percent
to 2 dp like the payload's.

  uv run --no-project --with openpyxl python scripts/chow-2023/build.py
"""

import hashlib
import json
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
WORKBOOK = HERE / "2023_office_of_the_mayor.xlsx"
OUT = HERE.parent.parent / "src" / "data" / "chow-2023.json"
SHA256 = "8077b5b94ebd25fb97d34cdeaeeadfb8a8c058d4530c06ab1ad56665a3fe435a"
SOURCE = "https://open.toronto.ca/dataset/elections-official-by-election-results/"

# Chow's 2026 candidacy id, as the payload and the canonical results carry it: the page finds her
# by it, never by name (another Chow is on the 2026 ballot).
CANDIDACY_ID_2026 = "can_24054bba2bed56dab41293ec7d138d41"
BALLOT_NAME_2023 = "Chow Olivia"
# The certified citywide result, as the canonical results carry it.
CITY_VOTES, CITY_VALID = 269_372, 724_638


def ward_totals(ws) -> tuple[str, dict[str, int]]:
    """The sheet's ward number and each candidate's ward total, checked against its totals row."""
    rows = [r for r in ws.iter_rows(values_only=True) if r and r[0] is not None]
    title, header = rows[0][0], rows[1]
    assert header[0] == "Subdivision" and header[-1] == "Total", (ws.title, header)
    num = title.split()[2]
    assert ws.title == f"Ward {num}", (ws.title, title)
    candidates = {r[0]: r[-1] for r in rows[3:] if not r[0].endswith("Totals")}
    totals = next(r for r in rows if r[0] == f"City Ward {num} Totals")
    assert sum(candidates.values()) == totals[-1], (
        num,
        sum(candidates.values()),
        totals[-1],
    )
    for r in rows[3:]:
        if not r[0].endswith("Totals"):
            assert sum(r[1:-1]) == r[-1], (num, r[0])
    return num, candidates


def main() -> None:
    assert hashlib.sha256(WORKBOOK.read_bytes()).hexdigest() == SHA256, (
        "workbook changed"
    )
    colours = json.loads(
        (HERE.parent / "candidate-colours" / "candidates.json").read_text()
    )
    assert any(
        c["id"] == CANDIDACY_ID_2026 and c["name"] == "Olivia Chow" for c in colours
    )

    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    wards = {}
    for name in wb.sheetnames[1:]:
        num, candidates = ward_totals(wb[name])
        assert len(candidates) == 102, (num, len(candidates))
        chow, valid = candidates[BALLOT_NAME_2023], sum(candidates.values())
        wards[num] = {
            "votes": chow,
            "valid": valid,
            "share": round(100 * chow / valid, 2),
        }
    assert list(wards) == [str(n) for n in range(1, 26)], list(wards)
    assert sum(w["votes"] for w in wards.values()) == CITY_VOTES
    assert sum(w["valid"] for w in wards.values()) == CITY_VALID

    out = {
        "candidacy_id": CANDIDACY_ID_2026,
        "source": "City of Toronto, certified results of the 2023 by-election for mayor",
        "source_url": SOURCE,
        "licence": "Contains information licensed under the Open Government Licence – Toronto.",
        "citywide": {
            "votes": CITY_VOTES,
            "valid": CITY_VALID,
            "share": round(100 * CITY_VOTES / CITY_VALID, 2),
        },
        "wards": wards,
    }
    OUT.write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n")
    print(
        f"wrote {OUT.relative_to(HERE.parent.parent)}: citywide {out['citywide']['share']}%"
    )


if __name__ == "__main__":
    main()
